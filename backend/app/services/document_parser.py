import fitz  # PyMuPDF
from docx import Document
from pptx import Presentation
import pytesseract
from PIL import Image
import io
from pathlib import Path
from typing import List, Dict, Optional

class DocumentParser:
    """文档解析服务：支持PDF/Word/PPT"""
    
    @staticmethod
    async def parse_pdf(file_path: str) -> Dict:
        """解析PDF文件"""
        chunks = []
        metadata = {"title": "", "pages": 0, "has_images": False}
        
        try:
            doc = fitz.open(file_path)
            metadata["pages"] = len(doc)
            metadata["title"] = doc.metadata.get("title", Path(file_path).stem)
            
            for page_num, page in enumerate(doc):
                text = page.get_text()
                if text.strip():
                    chunks.append({"index": page_num, "content": text.strip(), "type": "text"})
                
                # OCR处理图片
                images = page.get_images()
                if images:
                    metadata["has_images"] = True
                    for img_index, img in enumerate(images):
                        try:
                            xref = img[0]
                            pix = fitz.Pixmap(doc, xref)
                            if pix.n >= 5:
                                pix = fitz.Pixmap(fitz.csRGB, pix)
                            img_data = pix.tobytes("png")
                            pil_img = Image.open(io.BytesIO(img_data))
                            ocr_text = pytesseract.image_to_string(pil_img, lang='chi_sim+eng')
                            if ocr_text.strip():
                                chunks.append({"index": f"{page_num}_img_{img_index}", "content": ocr_text.strip(), "type": "ocr"})
                        except Exception:
                            continue
            doc.close()
        except Exception as e:
            return {"error": str(e), "chunks": [], "metadata": metadata}
        
        return {"chunks": chunks, "metadata": metadata}
    
    @staticmethod
    async def parse_docx(file_path: str) -> Dict:
        """解析Word文档"""
        chunks = []
        metadata = {"title": Path(file_path).stem, "paragraphs": 0}
        
        try:
            doc = Document(file_path)
            for i, para in enumerate(doc.paragraphs):
                if para.text.strip():
                    chunks.append({"index": i, "content": para.text.strip(), "type": "paragraph"})
            metadata["paragraphs"] = len(chunks)
            
            # 解析表格
            for table_idx, table in enumerate(doc.tables):
                table_text = []
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells]
                    table_text.append(" | ".join(row_text))
                if table_text:
                    chunks.append({"index": f"table_{table_idx}", "content": "\n".join(table_text), "type": "table"})
        except Exception as e:
            return {"error": str(e), "chunks": [], "metadata": metadata}
        
        return {"chunks": chunks, "metadata": metadata}
    
    @staticmethod
    async def parse_pptx(file_path: str) -> Dict:
        """解析PPT文件"""
        chunks = []
        metadata = {"title": Path(file_path).stem, "slides": 0}
        
        try:
            prs = Presentation(file_path)
            metadata["slides"] = len(prs.slides)
            
            for slide_num, slide in enumerate(prs.slides):
                slide_text = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_text.append(shape.text.strip())
                if slide_text:
                    chunks.append({"index": slide_num, "content": "\n".join(slide_text), "type": "slide"})
        except Exception as e:
            return {"error": str(e), "chunks": [], "metadata": metadata}
        
        return {"chunks": chunks, "metadata": metadata}

    @staticmethod
    async def parse_txt(file_path: str) -> Dict:
        """解析纯文本文件"""
        chunks = []
        metadata = {"title": Path(file_path).stem, "lines": 0}

        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()

            # 按段落分割
            paragraphs = [p.strip() for p in content.split('\n\n') if p.strip()]
            for i, para in enumerate(paragraphs):
                chunks.append({"index": i, "content": para, "type": "paragraph"})
            metadata["lines"] = len(paragraphs)
        except Exception as e:
            return {"error": str(e), "chunks": [], "metadata": metadata}

        return {"chunks": chunks, "metadata": metadata}

